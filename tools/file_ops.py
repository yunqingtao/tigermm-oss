TOOL = {
    "name": "file_ops",
    "description": "File operations: read, write, list, delete, copy, move files and directories.",
    "keywords": ["读取", "写入", "文件", "目录", "列出", "删除", "复制", "移动", "read", "write", "file", "list"],
    "params": [
        {"name": "action", "type": "str", "required": True,
         "enum": ["read", "write", "list", "delete", "copy", "move", "mkdir", "search", "walk"],
         "description": "Operation to perform"},
        {"name": "path", "type": "str", "required": True,
         "description": "File or directory path"},
        {"name": "content", "type": "str", "required": False,
         "description": "Content to write (for 'write')"},
        {"name": "dest", "type": "str", "required": False,
         "description": "Destination path (for copy/move)"},
        {"name": "offset", "type": "int", "required": False,
         "description": "Line number to start reading from"},
        {"name": "limit", "type": "int", "required": False,
         "description": "Maximum lines to read"},
    ]
}
def _safe_int(val, default=0):
    try:
        return int(val) if val != "" and val is not None else default
    except (ValueError, TypeError):
        return default


def _perm_msg(e) -> str:
    """越权/权限错误文案 —— 保留真因, 不叠加前缀。

    2026-09-23: 原写法 f"Access denied: {e}", 而 e 往往已是
    "Access denied: <路径>"(_resolve 里包的) -> 用户看到
    "Access denied: Access denied: <路径>", 真因(为什么被拒)全丢。
    故: e 里已含 denied/outside 就直接用, 不再套前缀。
    """
    _m = str(e)
    if "denied" in _m.lower() or "outside" in _m.lower():
        return _m
    return f"Access denied: {_m}"

"""
File operations plugin v2.0 — full Hermes-level file capabilities.
read(paginated) / write / list / walk(recursive+glob) / search(grep)
/ copy / move / delete / rename / patch / mkdir / exists / stat
All paths validated through storage.file_secure.
"""
import os as _os, json, sys, fnmatch, shutil, re as _re
from pathlib import Path

# ── 路径参数强校验 (2026-08-20 硬防线): 中文说明文字混入 → 拒绝 ──
_CJK_RE = _re.compile(r'[\u4e00-\u9fff]')

# 2026-08-20 v2: 中文指令词前缀 — 模型把 "把时间写入agent_test.txt" 整个当 path
# (相对路径无盘符, 原校验不拦 → 创建垃圾文件 "把时间写入agent_test.txt" 于项目根)
_CMD_PREFIX_RE = _re.compile(
    r'^(把|将|写|写入|保存|存|创建|新建|读取|读|打开|获取|添加|追加|覆盖|更新|修改|'
    r'内容|文字|文本|信息|数据|时间|日期|一行|为|叫|名为|名字|称|在|到|请|帮我|我要|'
    r'你好世界|你好|世界|本地|测试|完成)\s*[^A-Za-z:.\\\\/]*', _re.DOTALL)

# 合法相对路径前缀 (safe_path 支持): 桌面/文档/下载/盘符/./../
_LEGAL_REL_PREFIX = _re.compile(r'^(?:desktop|桌面|documents|文档|downloads|下载|D盘|C盘|[A-Za-z]:|[.][./\\\\]?)[\\\\/]?')


def _validate_path_arg(p, arg_name="path"):
    """路径参数强校验 (2026-08-20 硬防线):
    1. 剥离引号
    2. 绝对路径段后混入中文说明文字(如 'C:\\...\\a.txt 帮我保存') → 拒绝执行
    3. 相对路径以中文指令词开头(如 '把时间写入agent_test.txt'/'把你好世界写入x.txt') → 拒绝
    4. 合法中文文件名/目录(路径段内部中文) → 放行
    """
    if not isinstance(p, str) or not p.strip():
        return p
    p0 = p.strip().strip('"').strip("'")
    # ★★ 2026-09-23 加: **路径真实存在 → 直接放行**, 不再猜"是不是说明文字"。
    #   实测: `D:\tmm-考卷\题目\A1 查询.txt` 因**文件名里有空格**,
    #   被下面那条"按空白切分"的正则把 ` 查询.txt` 判成"路径后附中文说明"而拒读
    #   (用户从题目目录粘路径 → 读不了)。
    #   "真实存在"是最强判据: 存在的路径不可能是模型编的描述文字。
    try:
        if _os.path.exists(p0):
            return p0
    except Exception:
        pass
    m = _re.match(r'^([A-Za-z]:[\\/][^\s，。；、！？"\'<>|?*]*)(.*)$', p0, _re.DOTALL)
    if m:
        path_part, tail = m.group(1), m.group(2)
        if tail and _CJK_RE.search(tail):
            raise ValueError(
                f"{arg_name} 参数包含中文说明文字「{tail.strip()[:20]}」— 只允许纯文件路径, "
                f"禁止在路径后附加描述。请重新调用工具, 只传路径本身。")
        # 文件名段: 扩展名之后还有中文 → 说明文字混入 (如 'a.txt的内容')
        _last_seg = path_part.rsplit('\\', 1)[-1].rsplit('/', 1)[-1]
        if '.' in _last_seg:
            _after_ext = _last_seg[_last_seg.rfind('.') + 1:]
            if _after_ext and _CJK_RE.search(_after_ext):
                raise ValueError(
                    f"{arg_name} 参数中「{_last_seg[:30]}」的文件扩展名后混入中文说明文字 — "
                    f"只允许纯文件路径, 请重新调用工具。")
        if tail:
            # 路径段后还有内容但无中文(如含空格的路径后续) → 保留原始串
            return p0
        return path_part
    # ── 相对路径: 中文指令词污染检测 (2026-08-20 v2) ──
    # "把时间写入agent_test.txt" → 以指令词开头 → 拒绝 (模型把指令当路径)
    if _CMD_PREFIX_RE.match(p0) and not _LEGAL_REL_PREFIX.match(p0):
        # 剥指令词后剩余
        _stripped = _CMD_PREFIX_RE.sub('', p0).strip()
        # 如果剩余是纯文件名(含扩展名) → 明确拒绝, 提示模型单独传 path/content
        if _stripped and '.' in _stripped:
            raise ValueError(
                f"{arg_name} 参数「{p0[:40]}」以中文指令词开头 — 模型把指令与文件名混在一起了。"
                f"请重新调用: path 只传纯文件路径(如 'agent_test.txt'), 内容单独放 content 参数。")
        # 剩余无文件名特征 → 纯指令, 同样拒绝
        if _CJK_RE.search(_stripped) and not _re.search(r'\.[A-Za-z0-9]{1,5}$', _stripped):
            raise ValueError(
                f"{arg_name} 参数「{p0[:40]}」是中文指令而非文件路径 — 拒绝执行。"
                f"请重新调用工具, path 传纯文件路径。")
    return p0

PLUGIN = {
    "name": "file_ops",
    "description": "File ops: read/write/list/walk/search/copy/move/delete/rename/patch/mkdir/stat",
    "version": "2.0",
    "requires": [],
    "trigger": ["读文件", "写文件", "列出", "ls", "dir", "文件", "read", "write", "list",
                "读取", "遍历", "扫描", "找文件", "搜索文件", "walk", "scan",
                "复制", "拷贝", "copy", "移动", "mv", "move",
                "删除", "del", "rm", "rename", "重命名",
                "grep", "search", "查找", "搜索", "替换", "patch"],
    "permission": ["file_read", "file_write"],
    "category": "data",
}

async def run(**kwargs) -> dict:
    action = kwargs.get("action", "read")

    # ── 路径参数强校验 (2026-08-20 硬防线): 中文说明文字混入 → 拒绝 ──
    try:
        raw_path = _validate_path_arg(kwargs.get("path", ""), "path")
        dest = _validate_path_arg(kwargs.get("dest", kwargs.get("to", "")), "dest")
    except ValueError as _ve:
        return {"success": False, "error": f"参数校验失败: {_ve}"}
    content = kwargs.get("content", "")
    # ── ★ 2026-09-22 修 (用户实测报错 "Is a directory: ."): 写盘必须有**具体文件名** ──
    #   原来 path 缺省值是 "." ⇒ "写一首诗" 这类没给文件名的请求会真去写当前目录,
    #   报出用户看不懂的 `Is a directory: .`。缺文件名 → 诚实说缺什么, 不拿 "." 去撞。
    if action in ("write",) and (not raw_path.strip() or raw_path.strip() in (".", "./", ".\\")):
        return {"success": False,
                # ★ 消息里必须保留英文 "directory": 既有回归测试
                #   tests/test_cli_regression.py::test_dot_rejected_for_write 断言 error 含 "directory"
                "error": "没有指定文件名 —— path 为空或指向一个目录 (directory)。"
                         "请给出具体文件, 例如: 桌面/诗.txt"}

    # ── write 空内容防线 (2026-08-20 v2): 模型把 content 揉进 path 导致 content 空 → 拒绝 ──
    if action in ("write",) and not str(content or "").strip():
        return {"success": False,
                "error": "写入内容为空: content 参数缺失或为空。模型可能把内容误放进了 path — "
                         "请重新调用: path 传纯文件路径, 内容放 content 参数。"}
    pattern = kwargs.get("pattern", kwargs.get("query", ""))
    old_str = kwargs.get("old_string", kwargs.get("old", ""))
    new_str = kwargs.get("new_string", kwargs.get("new", ""))
    offset = _safe_int(kwargs.get("offset", 0))
    limit = _safe_int(kwargs.get("limit", 500))
    glob_pat = kwargs.get("glob", kwargs.get("file_glob", "*"))
    max_depth = _safe_int(kwargs.get("max_depth", kwargs.get("depth", 10)))

    # Path validation
    try:
        from storage.file_secure import safe_path, safe_read, safe_write, safe_list_dir
        from config.settings import PROJECT_ROOT as _prj_root
    except ImportError as e:
        return {"success": False, "error": f"Security module unavailable: {e}"}

    def _resolve(p, must_exist=False):
        """Resolve path: absolute paths pass through, relative paths resolve to PROJECT_ROOT.
        2026-08-20 v3: 裸文件名(无盘符/无目录分隔符) → 默认桌面 (用户语义"写文件"=桌面)。
        带目录的路径(core/pipeline.py) 保持项目根。"""
        _p_original = p
        if isinstance(p, str):
            # Strip quotes that models sometimes wrap paths in
            p = p.strip().strip('"').strip("'")
            # Absolute Windows path? Use directly (after security check)
            if _os.path.isabs(p):
                _denied = False
                try:
                    target = safe_path(p, _prj_root)
                except PermissionError:
                    # 路径越权 (不在允许根内) ≠ 文件不存在。旧代码一律吞成 None
                    # → 用户看到 "File not found" 会去查文件, 真因其实是路径不允许。
                    target, _denied = None, True
                if target and (not must_exist or target.exists()):
                    # 绝对路径的 write 目标若落在目录上 → 走 L186 的友好报错, 别提前返回目录
                    # (提前返回会让写入直接撞 PermissionError, 报成 "Access denied" 之类看不懂的错)
                    if action == "write" and target.is_dir():
                        raise IsADirectoryError(
                            f"Is a directory: {_p_original}. Specify a filename.")
                    return target
                if must_exist:
                    if _denied:
                        raise PermissionError(
                            f"Access denied (outside allowed roots): {_p_original}")
                    raise FileNotFoundError(f"File not found: {_p_original} (resolved to {p})")
            # 裸文件名 (无盘符, 无 / \ 分隔符, 非 . / .. / ./ 前缀) → 桌面 (2026-08-20 v3)
            # ./ 或 ..\ 前缀 = 明确"当前目录" → 保持项目根
            _norm = _os.path.normpath(p)
            _bare = _norm.replace('\\', '/')
            _p_lower = p.strip().lower()
            if (_bare not in ('.', '..') and '/' not in _bare
                    and not _p_lower.startswith(('./', '.\\', '../', '..\\'))
                    and not _bare.lower().startswith(('desktop', '桌面', 'documents', '文档',
                                                       'downloads', '下载', 'd盘', 'c盘'))
                    and not _os.path.splitdrive(_norm)[0]):
                # 2026-09-18 修正: 只有"文件名"(带扩展名)才映射到桌面。
                # 裸目录名 (core / tools) 是项目内相对路径 —— 之前一律映射桌面,
                # 导致 "列出 core 目录的文件" 去找 ~/Desktop/core 而失败 (既有 bug)。
                #   agent_test.txt (有扩展名)      → 桌面  (保留 2026-08-20 的修法)
                #   core (无扩展名, 项目内存在)     → 项目根
                #   报告 (无扩展名, 写意图, 项目内无) → 桌面
                _has_ext = bool(_os.path.splitext(_norm)[1])
                _in_project = _os.path.exists(_os.path.join(str(_prj_root), _norm))
                if _has_ext or (action == "write" and not _in_project):
                    _home_desktop = _os.path.join(_os.path.expanduser("~"), "Desktop")
                    p = _os.path.join(_home_desktop, _norm)
            # Relative path
            p = _os.path.normpath(p)
            p = p.removeprefix('.\\\\').removeprefix('./').removeprefix('\\\\').removeprefix('/')
        try:
            target = safe_path(p, _prj_root)
        except PermissionError:
            target = None
        if target and target.is_dir() and action == "write":
            # Reject bare directory references — model shouldn't write to a dir
            _bare = raw_path.strip().rstrip('\\/')
            if _bare in ('.', '..', '') or _os.path.isdir(raw_path.strip()) and not _os.path.splitext(raw_path.strip())[1]:
                raise IsADirectoryError(f"Is a directory: {raw_path}. Specify a filename.")
            # Auto-filename for directory write targets
            import datetime
            ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            p = _os.path.join(p, f"doc_{ts}.txt")
        try:
            t = safe_path(p, _prj_root)
            if must_exist and not t.exists():
                raise FileNotFoundError(f"File not found: {_p_original} (resolved to {t})")
            return t
        except PermissionError as e:
            # 2026-09-23: 保留真因 (原写法用裸路径覆盖消息 -> 用户只看到路径, 查不出为什么被拒)
            _why = str(e)
            raise PermissionError(_why if ("denied" in _why.lower() or "outside" in _why.lower())
                                  else f"Access denied: {p}") from e

    # ═══════════════════════════════════
    # READ (with offset/limit pagination)
    # ═══════════════════════════════════
    if action == "read":
        try:
            target = _resolve(raw_path, must_exist=True)
            if target.is_dir():
                # Auto-redirect: reading a directory → list it
                return await run(action="list", path=raw_path)
            # Excel (.xlsx): parse with openpyxl into readable text
            if target.suffix.lower() in (".xlsx", ".xlsm"):
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(target, read_only=True, data_only=True)
                    sheet_names = list(wb.sheetnames)
                    parts = []
                    for ws in wb.worksheets:
                        rows = []
                        for row in ws.iter_rows(values_only=True):
                            cells = [str(c) for c in row if c is not None]
                            if cells:
                                rows.append("  " + " | ".join(cells))
                        if rows:
                            parts.append(f"[Sheet: {ws.title}]")
                            parts.extend(rows[:200])
                    wb.close()
                    if not parts:
                        return {"success": True, "output": f"{target.name}: 空工作簿 (0 数据行)"}
                    n_rows = len(parts) - len(sheet_names)
                    meta = f"{target.name}: {len(sheet_names)} sheet(s), 前 {min(200, n_rows)} 行"
                    return {"success": True, "output": "\n".join(parts[:500]), "meta": meta}
                except ImportError:
                    return {"success": False, "error": f"{target.name}: 读取 xlsx 需要 openpyxl (pip install openpyxl)"}
                except Exception as e:
                    return {"success": False, "error": f"{target.name}: xlsx 解析失败: {e}"}
            lines = target.read_text(encoding="utf-8").split("\n")
            total = len(lines)
            if offset > 0 or limit < total:
                chunk = lines[offset:offset + limit]
                preview = "\n".join(chunk)
                chunk_end = min(offset + limit, total)
                done = chunk_end >= total
                meta_str = f"Lines {offset+1}-{chunk_end} of {total}"
                if done:
                    meta_str += " [FILE COMPLETE - all lines read]"
                else:
                    meta_str += f" [MORE AVAILABLE - use offset={chunk_end} for next page]"
                return {"success": True, "output": preview,
                        "meta": meta_str,
                        "total_lines": total, "offset": offset, "limit": limit, "complete": done}
            return {"success": True, "output": "\n".join(lines),
                    "meta": f"{total} lines [FILE COMPLETE]",
                    "total_lines": total, "complete": True}
        except FileNotFoundError as e:
            return {"success": False, "error": str(e)}
        except PermissionError as e:
            return {"success": False, "error": _perm_msg(e)}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ═══════════════════
    # WRITE
    # ═══════════════════
    elif action == "write":
        try:
            target = _resolve(raw_path)
            # ── content 污染检测 (2026-08-20 v3): 模型把整句指令塞进 content ──
            # 例: content="桌面 文档 把你好世界 agent_test.txt" 含自身文件名 → 拒绝
            _c = str(content or "")
            _fname = _os.path.basename(str(target))
            if _fname and _fname in _c and len(_c) > len(_fname) + 4:
                return {"success": False,
                        "error": f"写入内容疑似污染: content 包含文件名「{_fname}」和多余文字 — "
                                 f"模型把指令/文件名混进了内容。请重新调用: path 传纯路径, content 只传要写入的纯文本内容。"}
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding="utf-8")
            # 2026-08-20 v3: 返回绝对路径, 模型才知道文件实际落点(裸文件名→桌面)
            return {"success": True, "output": f"Written {len(content)} chars to {target}"}
        except PermissionError as e:
            return {"success": False, "error": _perm_msg(e)}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ═══════════════════
    # LIST
    # ═══════════════════
    elif action == "list":
        try:
            target = _resolve(raw_path, must_exist=True)
            if not target.is_dir():
                return {"success": False, "error": f"Not a directory: {raw_path}"}
            items = sorted(target.iterdir())
            lines = []
            for item in items[:50]:
                tag = "/" if item.is_dir() else ""
                try:
                    sz = item.stat().st_size
                    lines.append(f"  {item.name}{tag} ({sz:,}B)")
                except Exception:
                    lines.append(f"  {item.name}{tag}")
            return {"success": True, "output": "\n".join(lines) or "(empty)"}
        except PermissionError as e:
            return {"success": False, "error": _perm_msg(e)}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ══════════════════════════════════
    # WALK (recursive + glob filter)
    # ══════════════════════════════════
    elif action == "walk":
        try:
            target = _resolve(raw_path, must_exist=True)
            if not target.is_dir():
                return {"success": False, "error": f"Not a directory: {raw_path}"}
            results = []
            for root, dirs, files in _os.walk(str(target)):
                depth = root.replace(str(target), "").count(_os.sep)
                if depth > max_depth:
                    dirs[:] = []
                    continue
                for f in files:
                    if fnmatch.fnmatch(f, glob_pat):
                        fpath = _os.path.join(root, f)
                        size = _os.path.getsize(fpath)
                        results.append(f"{fpath} ({size:,}B)")
                if len(results) >= 200:
                    results.append("... (truncated at 200 results)")
                    break
            return {"success": True, "output": "\n".join(results) or "(no matches)"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ══════════════════════════════════
    # SEARCH (grep in files)
    # ══════════════════════════════════
    elif action in ("search", "grep"):
        if not pattern:
            return {"success": False, "error": "Missing 'pattern' for search"}
        try:
            target = _resolve(raw_path, must_exist=True)
            results = []
            if target.is_file():
                with open(target, "r", encoding="utf-8", errors="replace") as f:
                    for i, line in enumerate(f, 1):
                        if pattern.lower() in line.lower():
                            results.append(f"{target.name}:{i}: {line.rstrip()[:200]}")
                            if len(results) >= 50:
                                break
            elif target.is_dir():
                for root, dirs, files in _os.walk(str(target)):
                    depth = root.replace(str(target), "").count(_os.sep)
                    if depth > max_depth:
                        dirs[:] = []
                        continue
                    for fname in files:
                        if fnmatch.fnmatch(fname, glob_pat):
                            fpath = _os.path.join(root, fname)
                            try:
                                with open(fpath, "r", encoding="utf-8", errors="replace") as f:
                                    for i, line in enumerate(f, 1):
                                        if pattern.lower() in line.lower():
                                            results.append(f"{fpath}:{i}: {line.rstrip()[:200]}")
                                            if len(results) >= 50:
                                                break
                            except Exception:
                                pass
                            if len(results) >= 50:
                                break
                    if len(results) >= 50:
                        break
            return {"success": True, "output": "\n".join(results) or f"No matches for '{pattern}'"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ═══════════════════
    # COPY
    # ═══════════════════
    elif action in ("copy", "cp"):
        if not dest:
            return {"success": False, "error": "Missing 'dest' for copy"}
        try:
            src = _resolve(raw_path, must_exist=True)
            dst = _resolve(dest) if _os.path.isabs(dest) else safe_path(dest, _prj_root)
            dst.parent.mkdir(parents=True, exist_ok=True)
            if src.is_dir():
                shutil.copytree(str(src), str(dst), dirs_exist_ok=True)
            else:
                shutil.copy2(str(src), str(dst))
            return {"success": True, "output": f"Copied {raw_path} -> {dest}"}
        except PermissionError as e:
            return {"success": False, "error": _perm_msg(e)}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ═══════════════════
    # MOVE / RENAME
    # ═══════════════════
    elif action in ("move", "mv", "rename"):
        if not dest:
            return {"success": False, "error": "Missing 'dest' for move/rename"}
        try:
            src = _resolve(raw_path, must_exist=True)
            dst = _resolve(dest) if _os.path.isabs(dest) else safe_path(dest, _prj_root)
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(src), str(dst))
            return {"success": True, "output": f"Moved {raw_path} -> {dest}"}
        except PermissionError as e:
            return {"success": False, "error": _perm_msg(e)}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ═══════════════════
    # DELETE
    # ═══════════════════
    elif action in ("delete", "del", "rm"):
        try:
            target = _resolve(raw_path, must_exist=True)
            # ★★ 2026-09-24 修 (真 bug, 实测): 原来无条件调 input() 要确认 ——
            #   网页版/中继(无 stdin, DETACHED 进程) 里 input() 直接抛
            #   `EOFError: EOF when reading a line`, 并且以
            #   "Step N (file_ops): EOF when reading a line" 糊到用户脸上 (用户实测)。
            #   修法: 先探 stdin 是否可交互; 不可交互 → **不猜、直接拒绝**并说清怎么删。
            #   (删除是不可逆操作, 无人确认时宁可不动。)
            _stdin_ok = False
            try:
                import sys as _sys_del
                _stdin_ok = bool(_sys_del.stdin) and _sys_del.stdin.isatty()
            except Exception:
                _stdin_ok = False
            if not _stdin_ok:
                return {"success": False,
                        "error": ("删除需要人工确认，但当前没有可交互终端 (网页版/远程)。"
                                  "为安全起见没有执行删除。请在命令行模式操作，"
                                  "或改用 file_ops 的其它动作。")}
            try:
                confirm = input(f"\n  ⚠ 确认删除 {raw_path}? (y/n) ").strip().lower()
            except (EOFError, OSError) as _ie:
                return {"success": False,
                        "error": f"删除确认失败(无输入): {type(_ie).__name__} — 没有执行删除。"}
            if confirm not in ('y', 'yes'):
                return {"success": False, "error": "Delete cancelled by user"}
            if target.is_dir():
                shutil.rmtree(str(target))
            else:
                target.unlink()
            return {"success": True, "output": f"Deleted {raw_path}"}
        except PermissionError as e:
            return {"success": False, "error": _perm_msg(e)}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ═══════════════════
    # PATCH (targeted replace)
    # ═══════════════════
    elif action == "patch":
        if not old_str:
            return {"success": False, "error": "Missing 'old_string' for patch"}
        try:
            target = _resolve(raw_path, must_exist=True)
            text = target.read_text(encoding="utf-8")
            if old_str not in text:
                return {"success": False, "error": f"old_string not found in {raw_path}"}
            new_text = text.replace(old_str, new_str, 1)  # Replace first occurrence
            target.write_text(new_text, encoding="utf-8")
            return {"success": True, "output": f"Patched {raw_path} ({len(old_str)}->{len(new_str)} chars)"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ═══════════════════
    # MKDIR
    # ═══════════════════
    elif action == "mkdir":
        try:
            target = _resolve(raw_path)
            target.mkdir(parents=True, exist_ok=True)
            return {"success": True, "output": f"Created directory {raw_path}"}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ═══════════════════
    # EXISTS
    # ═══════════════════
    elif action == "exists":
        try:
            target = _resolve(raw_path)
            info = f"exists={target.exists()}, "
            if target.exists():
                info += f"is_file={target.is_file()}, is_dir={target.is_dir()}, "
                info += f"size={target.stat().st_size}B" if target.is_file() else ""
            return {"success": True, "output": info}
        except Exception as e:
            return {"success": False, "error": str(e)}

    # ═══════════════════
    # STAT
    # ═══════════════════
    elif action == "stat":
        try:
            target = _resolve(raw_path, must_exist=True)
            st = target.stat()
            import datetime
            mtime = datetime.datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
            info = f"size={st.st_size}B, mtime={mtime}, is_file={target.is_file()}, is_dir={target.is_dir()}"
            return {"success": True, "output": info}
        except Exception as e:
            return {"success": False, "error": str(e)}

    return {"success": False, "error": f"Unknown action: {action}"}
