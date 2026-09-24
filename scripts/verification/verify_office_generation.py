"""办公文档生成 (Word / PPT) — 常驻验证器 (ad-hoc 留档, 非 pytest 套件)

背景: 之前 `tiger_office` 的 14 个 action **全是读取/合并/整理, 一个"生成"都没有**;
`python-docx` / `python-pptx` **都没装** (代码有读取逻辑但一跑就报 not installed)。
这正是"功能看着有、实际没有"的来源, 也是 WorkBuddy 最核心的能力缺口。

落地 (2026-09-18):
  tools/tiger_office.py   +write_word (markdown 风格 → .docx, 自动抽标题, 中文字体 eastAsia)
                          +write_ppt  (## 分页 → .pptx, 自动加封面)
                          +_default_out_path (不给路径→桌面, 与 file_ops 裸名约定一致)
  core/pipeline.py        技能 DAG 新增 `llm` 内建步骤 —— 技能能"生成内容"而不只搬运
                          `_resolve_tmpl` 支持 $message 与**字符串内插值**(原来只支持整值占位符)
  core/skill_loader.py    `_KNOWN_TOOLS` 加 `llm` (否则 DAG 判无效 → 技能整个被丢弃)
  core/*.py 13 处扩展名白名单 补 `pptx?` (原来只认 docx/xlsx → .pptx 路径提不出来, 会落桌面)
  tmm_skills/write-office-doc, make-slides  两个技能 (SKILL.md, **零 pipeline 代码**)

★ 验证器对 llm 步骤**打桩**(确定性、零成本、可重复), 同时断言 prompt 真被拼对。
跑法:
    python -B scripts/verification/verify_office_generation.py
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent.parent

def _fill(s: str) -> str:
    """把探针模板里的 @@PROJ@@ 换成真实项目根。

    ★ 2026-09-21: 探针脚本写在临时目录里, 自己推不出项目根 (`__file__` 指向 %TEMP%)。
      原来是**写死绝对路径** —— 公开/换机就会指向别处。现在用占位符, 写入时填充。
    """
    # ★ 用**正斜杠**注入: 探针里的路径常被拼进中文话术再交给引擎解析, 反斜杠会与 "/"
    #   混用导致解析失败 (实测: 会当成"没给目录"而落到桌面)。正斜杠在 Windows 上同样合法。
    return s.replace("@@PROJ@@", str(ROOT).replace("\\", "/"))


PY = sys.executable

P, F = [], []
def chk(name, cond, detail=""):
    (P if cond else F).append(name)
    print(("PASS " if cond else "FAIL ") + name + (f"  <- {detail}" if detail and not cond else ""))


PROBE = r'''
import sys, os, json, re, asyncio, logging
sys.path.insert(0, r"@@PROJ@@"); os.chdir(r"@@PROJ@@")
logging.disable(logging.CRITICAL)
from main import _load_config
from core.model_client import ModelClient
from core.pipeline import Level4Pipeline
from core.knowledge import KnowledgeEngine
from core import skill_loader as SL
from config.settings import DATA_DIR

cfg = _load_config()
pl = Level4Pipeline(ModelClient(cfg), cfg)
pl._ke = KnowledgeEngine(DATA_DIR)
eng = SL.get_skill_engine()
# ★ 备份用户 mode.json (验证禁污染用户状态)
_MJ0 = r"@@PROJ@@/data/mode.json"
if os.path.exists(_MJ0) and not os.path.exists(_MJ0 + ".vfybak"):
    import shutil as _sh0
    _sh0.copy2(_MJ0, _MJ0 + ".vfybak")

SBX = r"@@PROJ@@/tmp/_office_verify"
os.makedirs(SBX, exist_ok=True)
out = {}

# ── 打桩 llm: 确定性、零成本; 同时捕获 prompt 供断言 ──
captured = {}
async def fake_generate(model, messages, temperature=0.7, max_tokens=4096, tools=None, think=None):
    captured["model"] = model
    captured["messages"] = messages
    captured["temperature"] = temperature
    # 按技能要求返回 markdown 文案 (与真模型同格式)
    if "幻灯片" in (messages[-1]["content"] or ""):
        return {"text": "# 项目进展汇报\n## 概览\n- 整体按计划推进\n- 核心指标达标\n## 小结\n- 风险可控", "error": None}
    return {"text": "# 本周工作周报\n## 一、完成\n- 修复 MCP 泄漏\n- 技能 DAG 接线\n## 二、计划\n- 办公文档生成", "error": None}
pl.model_client.generate = fake_generate

# 0. 技能 DAG 有效
out["dags"] = {}
for n in ["write-office-doc", "make-slides"]:
    d = eng.get_dag(n)
    out["dags"][n] = {"valid": bool(d and d.is_valid()),
                      "errors": (d.parse_errors if d else ["no dag"]),
                      "steps": len(d.steps) if d else 0}

# 1. 匹配
out["match"] = {}
for m in ["写周报，本周修了 MCP 泄漏", "写会议纪要", "做PPT讲项目进展", "做个汇报",
          "今天天气怎么样", "画个架构图"]:
    h = pl._skill_match(m)
    out["match"][m] = h[0] if h else None

# 2. $message 内插值 + 整值占位符
out["tmpl"] = {
    "interp": pl._resolve_tmpl("原始要求：$message", {}, {}, "写周报"),
    "whole": pl._resolve_tmpl("$params.p", {"p": "X"}, {}, "m"),
}
dag = eng.get_dag("write-office-doc")
out["prompt_has_message"] = "写周报" in pl._resolve_tmpl(
    dag.steps[0]["input"]["prompt"], {}, {}, "写周报")

async def _behave():
    # 3. ★ Word: 指定路径 + 抽标题不重复
    r = await pl._skill_dag_exec(f"写周报，本周修了 MCP 泄漏，存到 {SBX}/w.docx", 0.0)
    out["word"] = {"resp": str((r or {}).get("response"))[:130],
                   "intent": (r or {}).get("intent"), "skill": (r or {}).get("skill"),
                   "exists": os.path.exists(f"{SBX}/w.docx")}
    if out["word"]["exists"]:
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(f"{SBX}/w.docx")
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        out["word"]["paras"] = paras[:4]
        out["word"]["dup_title"] = len(paras) >= 2 and paras[0] == paras[1]
        rpr = doc.styles["Normal"].element.find(qn("w:rPr"))
        rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
        out["word"]["cjk"] = rf.get(qn("w:eastAsia")) if rf is not None else None
    out["llm_model_seen"] = captured.get("model")

    # 4. ★ PPT: 指定路径 (pptx 白名单修复) + 封面标题
    r2 = await pl._skill_dag_exec(f"做PPT讲项目进展，存到 {SBX}/p.pptx", 0.0)
    out["ppt"] = {"resp": str((r2 or {}).get("response"))[:130],
                  "exists": os.path.exists(f"{SBX}/p.pptx")}
    if out["ppt"]["exists"]:
        from pptx import Presentation
        prs = Presentation(f"{SBX}/p.pptx")
        out["ppt"]["slides"] = len(prs.slides)
        cover = [sh.text_frame.text.strip() for sh in prs.slides[0].shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        out["ppt"]["cover"] = cover

    # 5. 不给路径 → 桌面
    r3 = await pl._skill_dag_exec("写一份简短的工作总结", 0.0)
    resp3 = str((r3 or {}).get("response"))
    out["desktop"] = {"resp": resp3[:140], "to_desktop": "Desktop" in resp3}
    dm = re.search(r"([A-Za-z]:\\[^\s]+\.docx)", resp3)
    if dm and os.path.exists(dm.group(1)):
        os.remove(dm.group(1))          # 清掉, 别在用户桌面留垃圾
        out["desktop"]["cleaned"] = True

    # 6. ★ 路由优先级 + 生成守卫 (2026-09-18: IR chain 的"写X到Y"曾把生成请求写成垃圾文件)
    import re as _re
    ROUTE = [
        ("写周报，本周修了 MCP 泄漏，存到 " + SBX + "/r.docx", "skill"),
        ("做PPT讲项目进展，存到 " + SBX + "/r.pptx",           "skill"),
        ("写一份简短的工作总结",                                "skill"),
        ("帮我写份会议纪要，讨论了发布节奏",                     "skill"),
        ("帮我看看这份周报",                                   "other"),   # 读, 不是生成
        ("总结一下今天的事",                                   "other"),
        ("写文件",                                          "other"),   # 无技能触发
        ("把桌面文件发给涛哥",                                 "other"),   # 既有链路不动
        ("发邮件给涛哥 主题:hi 内容:test",                      "other"),
    ]
    out["route"] = []
    for m, exp in ROUTE:
        try:
            rr = await pl.process(m, probe=True)
        except Exception as e:
            rr = {"intent": "EXC:" + type(e).__name__, "response": str(e)[:60]}
        out["route"].append({"msg": m, "exp": exp, "intent": rr.get("intent"),
                             "skill": rr.get("skill"), "resp": str(rr.get("response"))[:80]})

    # 6b. ★ 5 个新办公技能: DAG 合法性 + 分派正确性 (2026-09-18 扩展)
    NEW_SKILLS = ["write-proposal", "write-speech", "write-self-review", "write-resume",
                  "table-to-report"]
    from core import skill_loader as _SL
    _e = _SL.get_skill_engine()
    out["new_dags"] = {}
    for _n in NEW_SKILLS:
        _d = _e.get_dag(_n)
        out["new_dags"][_n] = {"valid": bool(_d and _d.is_valid()),
                               "errors": (_d.parse_errors if _d else ["no dag"]),
                               "steps": len(_d.steps) if _d else 0}
    DISPATCH = {
        "写个新功能上线的推广方案": "write-proposal",
        "写一份年会致辞，感谢团队": "write-speech",
        "写份述职报告，今年做了三个项目": "write-self-review",
        "帮我写份简历，五年测试经验": "write-resume",
        "把这个表格做成分析报告 " + SBX + "/x.xlsx": "table-to-report",
        "写周报，本周修了 MCP 泄漏": "write-office-doc",
    }
    out["dispatch"] = []
    for _m, _exp in DISPATCH.items():
        _h = pl._skill_match(_m, trigger_only=True)
        out["dispatch"].append({"msg": _m, "exp": _exp, "got": (_h[0] if _h else None)})

    # 6c. ★ 曾被更早分支劫持的两种消息形态 (现在必须归位到技能层)
    import openpyxl as _oxl
    _wb = _oxl.Workbook(); _ws = _wb.active
    _ws.append(["城市", "渠道", "销量"])
    for _r in [("北京", "线上", 120), ("广州", "线上", 200)]:
        _ws.append(list(_r))
    _XLSX = SBX + "/t.xlsx"; _wb.save(_XLSX); _wb.close()
    HIJACK = [
        # ① 长消息 + 文件路径 + 含 '分析'/'方案' (曾被 _analysis_kw 预读劫持)
        ("把这个表格做成分析报告 " + _XLSX + " 到 " + SBX + "/t.docx，要求包含数据概况、关键发现和结论建议",
         "table-to-report"),
        # ② 含 '先/再/最后' (曾被 TaskPlanner 当多步任务拆)
        ("先写公司背景，再写推广方案，最后给出预算和时间表，做成一份完整的项目方案",
         "write-proposal"),
    ]
    out["hijack"] = []
    for _m, _exp in HIJACK:
        try:
            _rr = await pl.process(_m, probe=True)
        except Exception as _ex:
            _rr = {"intent": "EXC:" + type(_ex).__name__, "response": str(_ex)[:60]}
        out["hijack"].append({"msg": _m[:50], "exp": _exp, "intent": _rr.get("intent"),
                              "skill": _rr.get("skill"), "resp": str(_rr.get("response"))[:80]})

    # 6d. ★ @模型 三态 (技能层前移到 Fast path 之前 → 必须自带"不抢用户选择"的保护)
    out["ext_model"] = []
    for _em, _want in (("deepseek", False), ("auto", True), ("ollama", True)):
        try:
            _r = await pl.process("写份周报，本周修了 MCP 泄漏", ext_model=_em, probe=True)
            _g = _r.get("intent") == "skill_dag"
        except Exception as _ex:
            _g = None
        out["ext_model"].append({"ext": _em, "want": _want, "got": _g})



    # 6e. ★ probe=True 必须零副作用 (data/mode.json 是用户状态)
    #     实测踩过: 验证流量把 mode 写成 plan + pending_text → 之后所有请求走提案路径,
    #     技能层永远轮不到 (本验证器当时 4 项 FAIL, intent=plan_proposal)。
    import hashlib as _hl
    _MJ = r"@@PROJ@@/data/mode.json"

    def _mh():
        return _hl.md5(open(_MJ, "rb").read()).hexdigest()
    out["modejsn"] = {}
    if os.path.exists(_MJ):
        open(_MJ, "w", encoding="utf-8").write(json.dumps(
            {"mode": "craft", "updated": 0, "pending": [], "pending_message": "",
             "pending_text": "", "pending_at": 0}, ensure_ascii=False, indent=2))
        _base = _mh()
        _rr = await pl.process("帮我做一份季度规划", mode_override="plan", probe=True)
        out["modejsn"] = {"probe_intent": _rr.get("intent"), "base": _base,
                           "after_probe": _mh()}

asyncio.run(_behave())

# 7. 生成守卫: 读类请求不得触发 llm 生成
out["guard"] = {
    "read_report": pl._skill_dag_exec is not None and not pl._looks_like_creation("帮我看看这份周报"),
    "write_report": pl._looks_like_creation("写份周报"),
}
# 清掉路由测试可能落桌面的文件
for pat in ("*周报*.docx", "工作总结*.docx", "会议纪要*.docx", "生成内容*.docx",
                    "文档_*.docx", "致辞*.docx", "述职*.docx", "分析*.docx"):
    for f in (__import__("pathlib").Path.home() / "Desktop").glob(pat):
        try: f.unlink()
        except Exception: pass

# ★ 还原 mode.json (探针自己改过 -- 必须留回原样)
_MJF = r"@@PROJ@@/data/mode.json"
if os.path.exists(_MJF):
    try:
        import shutil as _sh
        if os.path.exists(_MJF + ".vfybak"):
            _sh.move(_MJF + ".vfybak", _MJF)
    except Exception:
        pass

print("<<<J>>>" + json.dumps(out, ensure_ascii=False))
'''


def run_probe():
    fd, tmp = tempfile.mkstemp(suffix=".py"); os.close(fd)
    Path(tmp).write_text(_fill(PROBE), encoding="utf-8")
    try:
        r = subprocess.run([PY, "-B", tmp], capture_output=True, text=True,
                           encoding="utf-8", errors="replace", timeout=400, cwd=str(ROOT))
        d = json.loads(r.stdout.split("<<<J>>>")[1].strip().splitlines()[0])
        return d, r.stdout + r.stderr
    finally:
        Path(tmp).unlink(missing_ok=True)


def main():
    SBX = ROOT / "tmp" / "_office_verify"
    try:
        print("        —— 静态 ——")
        code = r'''
import importlib.util as iu
mods = {m: bool(iu.find_spec(m)) for m in ["docx", "pptx"]}
to = open(r"@@PROJ@@/tools/tiger_office.py", encoding="utf-8").read()
sl = open(r"@@PROJ@@/core/skill_loader.py", encoding="utf-8").read()
import json
print("<<<J>>>" + json.dumps({"mods": mods,
    "write_word": "_write_word" in to, "write_ppt": "_write_ppt" in to,
    "registered": '"write_word": _write_word' in to and '"write_ppt": _write_ppt' in to,
    "llm_known": '"llm"' in sl,
    "pptx_wl": "pptx?" in open(r"@@PROJ@@/core/pipeline.py", encoding="utf-8").read()}, ensure_ascii=False))
'''
        fd, tmp = tempfile.mkstemp(suffix=".py"); os.close(fd)
        Path(tmp).write_text(_fill(code), encoding="utf-8")
        try:
            r = subprocess.run([PY, "-B", tmp], capture_output=True, text=True,
                               encoding="utf-8", errors="replace", timeout=120, cwd=str(ROOT))
            st = json.loads(r.stdout.split("<<<J>>>")[1].strip().splitlines()[0])
        finally:
            Path(tmp).unlink(missing_ok=True)

        chk("python-docx 已装", st["mods"]["docx"] is True)
        chk("python-pptx 已装", st["mods"]["pptx"] is True)
        chk("write_word / write_ppt 已实现", st["write_word"] and st["write_ppt"])
        chk("两个 action 已注册", st["registered"] is True)
        chk("skill_loader 白名单含 llm", st["llm_known"] is True)
        chk("扩展名白名单含 pptx?", st["pptx_wl"] is True)
        for n in ["write-office-doc", "make-slides"]:
            f = ROOT / "tmm_skills" / n / "SKILL.md"
            chk(f"技能文件存在 {n}", f.exists())
            if f.exists():
                t = f.read_text(encoding="utf-8")
                chk(f"  {n} frontmatter 以 --- 开头", t.lstrip().startswith("---"))
                chk(f"  {n} 有 llm 步骤", "tool: llm" in t)

        print("        —— 行为 (llm 打桩, 确定性) ——")
        d, raw = run_probe()

        for n, v in d["dags"].items():
            chk(f"DAG 有效: {n}", v["valid"] is True, str(v["errors"])[:120])
        chk("匹配 写周报 → write-office-doc",
            d["match"]["写周报，本周修了 MCP 泄漏"] == "write-office-doc",
            str(d["match"]["写周报，本周修了 MCP 泄漏"]))
        chk("匹配 做PPT → make-slides", d["match"]["做PPT讲项目进展"] == "make-slides",
            str(d["match"]["做PPT讲项目进展"]))
        chk("匹配 天气 → None", d["match"]["今天天气怎么样"] is None)
        chk("匹配 架构图 → diagram (没被技能抢走)", d["match"]["画个架构图"] == "diagram",
            str(d["match"]["画个架构图"]))

        chk("$message 内插值生效", d["tmpl"]["interp"] == "原始要求：写周报", d["tmpl"]["interp"])
        chk("整值占位符保留原值", d["tmpl"]["whole"] == "X", str(d["tmpl"]["whole"]))
        chk("技能 prompt 里含用户原话(替换后)", d["prompt_has_message"] is True)
        chk("llm 步骤默认走 deepseek", d["llm_model_seen"] == "deepseek", str(d["llm_model_seen"]))

        w = d["word"]
        chk("★ Word: 指定路径生效", w["exists"] is True, w["resp"])
        chk("Word: intent=skill_dag / skill=write-office-doc",
            w["intent"] == "skill_dag" and w["skill"] == "write-office-doc", str(w)[:120])
        chk("★ Word: 标题不重复", w.get("dup_title") is False,
            str(w.get("paras"))[:110])
        chk("Word: 内容来自 llm 步骤", bool(w.get("paras")) and "周报" in (w["paras"][0] or ""),
            str(w.get("paras"))[:110])
        chk("Word: 中文字体写入 eastAsia", w.get("cjk") == "微软雅黑", str(w.get("cjk")))

        p = d["ppt"]
        chk("★ PPT: 指定路径生效 (pptx 白名单修复)", p["exists"] is True, p["resp"])
        chk("★ PPT: 封面标题来自 '# 标题' 行",
            bool(p.get("cover")) and "汇报" in (p["cover"][0] if p["cover"] else ""),
            str(p.get("cover")))
        chk("PPT: 页数 = 封面 + 2 页文案", p.get("slides") == 3, str(p.get("slides")))

        chk("★ 不给路径 → 落桌面", d["desktop"]["to_desktop"] is True, d["desktop"]["resp"])
        chk("桌面探测文件已清理", d["desktop"].get("cleaned") is True)

        print("        —— ★ 路由优先级 + 生成守卫 ——")
        rbad = [c for c in d["route"]
                if (c["exp"] == "skill") != (c["intent"] == "skill_dag")]
        for c in d["route"]:
            mark = "✓" if (c["exp"] == "skill") == (c["intent"] == "skill_dag") else "✗"
            print(f"          {mark} {c['msg'][:32]!r:36s} → {c['intent']}/{c.get('skill')}")
        chk("★ 生成类请求走技能层 (不被 IR chain 吞成写文件)", not rbad,
            str([c["msg"] for c in rbad]))
        chk("★ 读类请求不触发生成 (生成守卫)",
            d["guard"]["read_report"] is True and d["guard"]["write_report"] is True,
            str(d["guard"]))

        print("        —— ★ 新增 5 个办公技能 ——")
        for n, v in d["new_dags"].items():
            chk(f"DAG 有效: {n} (steps={v['steps']})", v["valid"] is True, str(v["errors"])[:110])
        dbad = [c["msg"] for c in d["dispatch"] if c["got"] != c["exp"]]
        for c in d["dispatch"]:
            print(f"          {'✓' if c['got'] == c['exp'] else '✗'} {c['msg'][:32]!r:36s} → {c['got']}")
        chk("★ 6 类生成请求分派正确 (专属技能优先)", not dbad, str(dbad))

        print("        —— ★ 曾被更早分支劫持的形态 (现须归位) ——")
        for c in d["hijack"]:
            ok = c["intent"] == "skill_dag" and c["skill"] == c["exp"]
            print(f"          {'✓' if ok else '✗'} {c['msg'][:40]!r:44s} → {c['intent']}/{c['skill']}")
        hbad = [c["msg"] for c in d["hijack"] if not (c["intent"] == "skill_dag" and c["skill"] == c["exp"])]
        chk("★ 长消息+路径(含'分析') 不再被 _analysis_kw 劫持", len(hbad) == 0, str(hbad))
        chk("★ 含'先/再'的生成请求 不再被 TaskPlanner 拆", len(hbad) == 0, str(hbad))

        print("        —— ★ probe=True 零副作用 (mode.json 不被测试改) ——")
        mj = d.get("modejsn") or {}
        chk("probe 流量确实走了 plan 提案路径", mj.get("probe_intent") == "plan_proposal", str(mj))
        chk("★ probe 后 mode.json 字节级不变", mj.get("after_probe") == mj.get("base"),
            f"base={str(mj.get('base'))[:12]} after={str(mj.get('after_probe'))[:12]}")

        print("        —— ★ probe 契约必须贯穿**所有入口** (web 层曾漏) ——")
        _ws = (ROOT / "web_server.py").read_text(encoding="utf-8")
        # web 的 /chat 支持请求体带 mode → 会写 mode.json。probe 流量必须不写。
        # ⚠ 必须定位到 **/chat 处理器内**的那一处 —— data.get("mode") 在 /mode 接口里
        #   也出现一次(那是切模式的正当入口, 不需要 probe 守卫), 直接 find 会取错窗口。
        _ci = _ws.find('@app.post("/chat")')
        _mi = _ws.find('data.get("mode")', _ci) if _ci != -1 else -1
        _mw = _ws[_mi:_mi + 400] if _mi != -1 else ""
        chk("★ web /chat 的 mode 写入有 probe 守卫", _mi != -1 and "not probe" in _mw,
            f"窗口: {_mw[:150]!r}")
        # ModeManager.set 只在真的变化时写盘 (否则前端每次请求都写磁盘)
        _mm = (ROOT / "core" / "modes.py").read_text(encoding="utf-8")
        _si = _mm.find("def set(self")
        # ★ 2026-09-23: 按**方法边界**切片, 不用固定字符数 ——
        #   原来写 _mm[_si:_si+700], 一旦 set() 里多几行注释/日志就把 _save() 挤出窗口
        #   ⇒ 假红 "没有只在变化时写盘" (实测: 加了个模式切换留痕日志就红了)。
        _end = _mm.find("\n    def ", _si + 1)
        _sw = _mm[_si:_end if _end != -1 else _si + 3000]
        _logical = [l.strip() for l in _sw.split("\n") if l.strip()]
        _save_inside = False
        _ind = None
        for l in _sw.split("\n"):
            if "if mode != self.state.get" in l:
                _ind = len(l) - len(l.lstrip())
            if "_save()" in l and _ind is not None:
                _save_inside = (len(l) - len(l.lstrip())) > _ind
        chk("★ ModeManager.set 仅在实际变化时写盘", _save_inside, f"{_logical[:8]}")

        print("        —— ★ 技能层位置: @模型 三态 (不抢用户显式选择) ——")
        embad = []
        for c in d["ext_model"]:
            if c["want"] != c["got"]:
                embad.append(c["ext"])
            print(f"          {'✓' if c['want'] == c['got'] else '✗'} @{c['ext']:10s} "
                  f"want_skill={c['want']!s:5s} got={c['got']!s:5s}")
        chk("★ @deepseek 不抢 / @auto 抢 / @ollama 抢", not embad, str(embad))

        # ★ 契约式 (2026-09-19 改): 原来按**行号**断言"技能块早于 4 个劫持者" ——
        #   技能层迁进 core/routes.py 后该断言必然失效(行号不存在了)。本意是
        #   "技能层必须赢过那些启发式", 现在由**优先级**保证 → 断言优先级契约更准。
        _src_txt = (ROOT / "core" / "pipeline.py").read_text(encoding="utf-8")
        chk("★ 技能层已迁入路由表 (process() 里不再有内联技能块)",
            "技能 DAG 层 · 强触发词优先" not in _src_txt)
        chk("process() 里不再有内联 TaskPlanner 块",
            "# TaskPlanner: complex multi-step tasks" not in _src_txt)
        chk("process() 里不再有内联 @模型 fast path",
            "Fast path: if user picked a model" not in _src_txt)

        # 静态解析 core/routes.py 的优先级常量 (免导入路径问题, 仍是契约式)
        _rt = (ROOT / "core" / "routes.py").read_text(encoding="utf-8")
        _P = {m.group(1): int(m.group(2)) for m in
              re.finditer(r"^(P_[A-Z_]+)\s*=\s*(\d+)", _rt, re.M)}
        chk("优先级常量齐备", len(_P) >= 8, str(_P))
        _want = [("技能 > TaskPlanner", "P_SKILL", "P_PLANNER"),
                 ("技能 > 分析预读", "P_SKILL", "P_ANALYSIS"),
                 ("技能 > 搜索", "P_SKILL", "P_SEARCH"),
                 ("技能 > IR chain", "P_SKILL", "P_IR_CHAIN"),
                 ("@模型 > 技能", "P_EXPLICIT_MODEL", "P_SKILL"),
                 ("命令 > @模型", "P_COMMAND", "P_EXPLICIT_MODEL")]
        for _nm, _hi, _lo in _want:
            if _hi in _P and _lo in _P:
                chk(f"★ 优先级契约: {_nm}", _P[_hi] > _P[_lo], f"{_P[_hi]} vs {_P[_lo]}")
            else:
                chk(f"★ 优先级契约: {_nm}", False, f"缺常量 {_hi}/{_lo}")

        print("        —— canonical ——")
        r = subprocess.run([PY, "-B", "-m", "pytest", "-q", "--no-header", "-p", "no:cacheprovider"],
                           cwd=str(ROOT), capture_output=True, text=True,
                           encoding="utf-8", errors="replace", timeout=580)
        tail = [l for l in r.stdout.split("\n") if "passed" in l or "failed" in l]
        print(f"        {tail[-1] if tail else r.stdout.strip()[-70:]}")
        chk("canonical 全绿 (0 failed)", r.returncode == 0 and "failed" not in r.stdout,
            r.stdout.strip()[-150:])
    finally:
        _bk = ROOT / "data" / "mode.json.vfybak"
        if _bk.exists():
            shutil.move(str(_bk), str(ROOT / "data" / "mode.json"))
        for dirp in (SBX, ROOT / "tmp" / "_office_smoke"):
            if dirp.exists():
                shutil.rmtree(dirp, ignore_errors=True)
        # 清桌面可能残留的探测文档
        for f in [x for pat in ("工作总结*.docx", "生成内容*.docx", "文档_*.docx",
                                 "*周报*.docx", "致辞*.docx", "述职*.docx")
                        for x in (Path.home() / "Desktop").glob(pat)]:
            f.unlink()

    print(f"\n结果: {len(P)} PASS / {len(F)} FAIL")
    for f in F:
        print("  -", f)
    return 1 if F else 0


if __name__ == "__main__":
    # ★ 快照/还原用户状态: 本验证器真跑 pipeline → data/perception.json 会被写 (实测)
    _perc = ROOT / "data" / "perception.json"
    _snap = _perc.read_bytes() if _perc.exists() else None
    try:
        _code = main()
    finally:
        if _snap is not None:
            _perc.write_bytes(_snap)
    sys.exit(_code)
